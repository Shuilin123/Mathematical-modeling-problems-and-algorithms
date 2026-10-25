# -*- coding: utf-8 -*-
"""数据读取、储能参数与结果文件写出。

约定：一天等分 144 个 10min 时段，第 t 个时段为 [10(t-1), 10t) 分钟；
附件 1/2/4 的第 t 列（列标签 0:10, 0:20, …, 0:00+1）即第 t 个时段的数值。
功率单位 kW，电价 元/kWh，时段长 DT = 1/6 h。
"""
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font

# ---------------- 路径 ----------------
ROOT = Path(__file__).resolve().parents[2]          # 项目根目录
DATA = ROOT / "附件" / "附件"
OUT = ROOT / "out"
FIG = ROOT / "fig"
for _p in (OUT, FIG):
    _p.mkdir(exist_ok=True)

# ---------------- 参数（附录 1） ----------------
DT = 1.0 / 6.0            # 时段长度 h
T = 144                   # 每日时段数
E_MAX = 12000.0           # 储能最大容量 kWh
P_MAX = 5000.0            # 最大充放电功率 kW
E_MIN = 1200.0            # 运行区间下限 kWh
E_CAP = 10800.0           # 运行区间上限 kWh
E_INIT = 6000.0           # 2025-01-01 0:00 储电量 kWh
ETA = 0.90                # 充/放电效率
K_ADD = 5.0               # 紧急购电电价倍率


# ---------------- 数据读取 ----------------
def load_day_profile():
    """附件1：单日电价、小区负载、光伏预测功率（各 144 点）。"""
    df = pd.read_excel(DATA / "附件1.xlsx")
    assert len(df) == T
    return (df["电价"].to_numpy(float), df["小区负载"].to_numpy(float),
            df["光伏发电预测功率"].to_numpy(float))


def load_year(sheet_file="附件2.xlsx"):
    """附件2：全年 365 天 x 144 时段的负载与光伏实际功率。"""
    xl = pd.ExcelFile(DATA / sheet_file)
    load = pd.read_excel(xl, sheet_name="小区负载").iloc[:, 1:].to_numpy(float)
    pv = pd.read_excel(xl, sheet_name="光伏发电实际功率").iloc[:, 1:].to_numpy(float)
    return load, pv


def load_price_year():
    """附件4：全年 365 天 x 144 时段的实时电价。"""
    return pd.read_excel(DATA / "附件4.xlsx").iloc[:, 1:].to_numpy(float)


def load_pv_forecast():
    """附件3：每天 0:00/6:00/12:00/18:00 发布的未来 24 个整点光伏预报 (365,4,24)。"""
    df = pd.read_excel(DATA / "附件3.xlsx")
    cols = ["预报%d小时" % i for i in range(1, 25)]
    return df[cols].to_numpy(float).reshape(365, 4, 24)


# ---------------- 结果文件写出 ----------------
BLOCKS = ["0:00-4:00", "4:00-8:00", "8:00-12:00",
          "12:00-16:00", "16:00-20:00", "20:00-24:00"]


def labels_144():
    """从附件5 模板读取 144 个时段标签。"""
    wb = openpyxl.load_workbook(DATA / "附件5" / "result1.xlsx")
    ws = wb["计划购电量"]
    return [ws.cell(r, 1).value for r in range(2, 146)]


def dates():
    """2025-02-01 — 2025-12-31 的 334 个日期字符串。"""
    return [str(np.datetime64("2025-01-01") + np.timedelta64(d, "D"))
            for d in range(31, 365)]


def _style(ws):
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")


def block_sum(x):
    """6 个 4 小时时段的合计 kWh。"""
    x = np.nan_to_num(np.asarray(x, float))
    return [x[24 * k:24 * (k + 1)].sum() * DT for k in range(6)]


def runs_of(add, min_kwh=0.5):
    """把逐时段紧急购电合并成连续时段 [(起, 止, kWh), …]。"""
    add = np.nan_to_num(np.asarray(add, float)) * DT
    out, t = [], 0
    while t < T:
        if add[t] > min_kwh:
            s = t
            while t < T and add[t] > min_kwh:
                t += 1
            out.append((s, t, add[s:t].sum()))
        else:
            t += 1
    return out


def seg_label(s, e):
    def hm(x):
        return "%d:%02d" % ((x * 10) // 60, (x * 10) % 60)
    return "%s-%s" % (hm(s), hm(e))


def write_result1(g, ch, dis, Q, price):
    """问题1 结果：144 段购电量 + 6 个 4h 时段充放电量与 0:00/24:00 储电量。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "计划购电量"
    ws.append(["时间段", "购电量"])
    for k, lab in enumerate(labels_144()):
        ws.append([lab, round(float(g[k] * DT), 4)])
    _style(ws)
    ws.column_dimensions["A"].width = 16

    ws2 = wb.create_sheet("充放电量")
    ws2.append(["时间段", "充电量", "放电量", "时刻", "储电量"])
    c, d = block_sum(ch), block_sum(dis)
    for k in range(6):
        row = [BLOCKS[k], round(c[k], 4), round(d[k], 4), None, None]
        if k == 0:
            row[3], row[4] = "0:00", round(float(Q[0]), 4)
        if k == 1:
            row[3], row[4] = "24:00", round(float(Q[144]), 4)
        ws2.append(row)
    _style(ws2)
    ws2.column_dimensions["A"].width = 16

    path = ROOT / "result1.xlsx"
    wb.save(path)
    print("已写出 %s：全天购电 %.2f kWh，购电费 %.2f 元"
          % (path.name, g.sum() * DT, (g * price * DT).sum()))


def write_multi(fname, plan, ch, dis, Q, add, adj=None):
    """问题2/3/4 结果：计划（+调整）购电量、充放电量、紧急购电量。"""
    ds = dates()
    lab = labels_144()
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "计划购电量"
    ws.append(["日期\\时间"] + lab)
    for i, d in enumerate(ds):
        ws.append([d] + [round(float(v * DT), 3) for v in plan[i]])
    _style(ws)
    ws.freeze_panes = "B2"

    if adj is not None:
        ws2 = wb.create_sheet("调整购电量")
        ws2.append(["日期\\时间"] + lab)
        for i, d in enumerate(ds):
            ws2.append([d] + [round(float(v * DT), 3) for v in adj[i]])
        _style(ws2)
        ws2.freeze_panes = "B2"

    ws3 = wb.create_sheet("充放电量")
    ws3.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
    for i, d in enumerate(ds):
        c, dd = block_sum(ch[i]), block_sum(dis[i])
        for k in range(6):
            row = [d if k == 0 else None, BLOCKS[k],
                   round(c[k], 3), round(dd[k], 3), None, None]
            if k == 0:
                row[4], row[5] = "0:00", round(float(Q[i][0]), 3)
            if k == 1:
                row[4], row[5] = "24:00", round(float(Q[i][144]), 3)
            ws3.append(row)
    _style(ws3)

    ws4 = wb.create_sheet("紧急购电量")
    ws4.append(["日期", "紧急购电时间段", "紧急购电量"])
    for i, d in enumerate(ds):
        rs = runs_of(add[i])
        if not rs:
            ws4.append([d, None, 0.0])
            continue
        for j, (s, e, e_kwh) in enumerate(rs[:4]):
            ws4.append([d if j == 0 else None, seg_label(s, e), round(float(e_kwh), 3)])
    _style(ws4)

    path = ROOT / fname
    wb.save(path)
    print("已写出 %s：计划购电 %.0f kWh，紧急购电 %.0f kWh"
          % (fname, plan.sum() * DT, add.sum() * DT))


if __name__ == "__main__":
    price, load, pv = load_day_profile()
    print("附件1 电价 %.4f~%.4f（均值 %.4f）" % (price.min(), price.max(), price.mean()))
    print("附件1 负载 %.1f~%.1f，光伏 %.1f~%.1f"
          % (load.min(), load.max(), pv.min(), pv.max()))
    L, P = load_year()
    print("附件2 %s，附件4 %s" % (L.shape, load_price_year().shape))
